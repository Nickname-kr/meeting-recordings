"""A4 meeting exports; generated in memory without a remote conversion service."""
from io import BytesIO
from pathlib import Path
from html import escape
import re
import unicodedata
import threading

SECTIONS = [('agenda', '안건'), ('notes', '논의 내용'), ('decisions', '결정 사항'), ('actions', '후속 업무')]
FONT_DIR = Path(__file__).parent / 'assets' / 'fonts'
FONT_LOCK = threading.Lock()


def clean(value):
    text = unicodedata.normalize('NFC', str(value or ''))
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', text).replace('\r\n', '\n').replace('\r', '\n')


def filename(record):
    title = re.sub(r'[\\/:*?"<>|\x00-\x1f]', '_', clean(record.get('title')))[:60].strip(' .') or '회의록'
    date = re.sub(r'[^0-9-]', '', str(record.get('date', '')))[:10]
    return f'{date}_{title}' if date else title


def metadata(record):
    return [('일시', f"{record.get('date', '')}  {record.get('start', '')} ~ {record.get('end', '')}"), ('장소', record.get('place', '')), ('참석자', ', '.join(record.get('people', []))), ('작성자', record.get('author', ''))]


def make_docx(record):
    from docx import Document
    from docx.shared import Mm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(22)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1']:
        style = doc.styles[name]
        style.font.name = 'NanumGothic'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'NanumGothic')
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
        for fonts in style.element.iter(qn('w:rFonts')):
            for attribute in list(fonts.attrib):
                if attribute.endswith('Theme'):
                    del fonts.attrib[attribute]
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)
    doc.styles['Title'].font.size = Pt(24)
    doc.styles['Title'].font.bold = True
    doc.styles['Subtitle'].font.size = Pt(14)
    doc.styles['Heading 1'].font.size = Pt(13)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(16)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(7)
    doc.add_paragraph('회의록', 'Title')
    doc.add_paragraph(clean(record.get('title')), 'Subtitle')
    for label, value in metadata(record):
        p = doc.add_paragraph()
        p.add_run(label + '  ').bold = True
        p.add_run(clean(value) or '—')
    for key, label in SECTIONS:
        doc.add_paragraph(label, 'Heading 1')
        for line in (clean(record.get(key)) or '—').split('\n'):
            p = doc.add_paragraph(line)
            p.paragraph_format.widow_control = True
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.core_properties.title = clean(record.get('title'))
    doc.core_properties.author = ''
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def make_pdf(record):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    with FONT_LOCK:
        for face, file in [('MeetingKR', 'NanumGothic-Regular.ttf'), ('MeetingKRBold', 'NanumGothic-Bold.ttf')]:
            if face not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(face, str(FONT_DIR / file)))
    body = ParagraphStyle('body', fontName='MeetingKR', fontSize=11, leading=17, wordWrap='CJK', spaceAfter=5)
    title = ParagraphStyle('title', parent=body, fontName='MeetingKRBold', fontSize=24, leading=32, spaceAfter=12, keepWithNext=True)
    subtitle = ParagraphStyle('subtitle', parent=body, fontSize=14, leading=21, spaceAfter=18, keepWithNext=True)
    heading = ParagraphStyle('heading', parent=body, fontName='MeetingKRBold', fontSize=13, leading=19, spaceBefore=16, spaceAfter=7, keepWithNext=True)
    def para(value, style=body):
        return Paragraph(escape(clean(value)).replace('\n', '<br/>') or '&#160;', style)
    story = [para('회의록', title), para(record.get('title'), subtitle)]
    for label, value in metadata(record):
        story.append(para(label + '  ' + (clean(value) or '—')))
    for key, label in SECTIONS:
        story.append(para(label, heading))
        for line in (clean(record.get(key)) or '—').split('\n'):
            story.append(para(line))
    output = BytesIO()
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('MeetingKR', 9)
        canvas.drawRightString(A4[0] - 22*mm, 12*mm, str(doc.page))
        canvas.restoreState()
    doc = SimpleDocTemplate(output, pagesize=A4, leftMargin=22*mm, rightMargin=22*mm, topMargin=20*mm, bottomMargin=20*mm, title=clean(record.get('title')), author='')
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()
